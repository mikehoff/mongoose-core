"""LPE1 stage 2 templates: the covering letter, the completed Form LPE1 pack, and the cover email.

Built from the first live release (Flat 7, 1 Example Road, 23 September 2026) so that rendering the Flat 7 job
reproduces the two PDFs that were sent, and any other flat gets the same documents with its own details.
See SPEC_stage2_LPE1_release_mongoose.md.

What lives where. This file holds layout and mechanics only. The wording, the answer set, the requester's
blank form, the enclosures and each job's facts are real data, so they live in the data directory, in
<data_dir>/rtm_templates/lpe1/:

  form_answers.json   every form value (text by field name, radios by export value) with {{placeholders}}
  form_overlays.json  text printed on the form after flattening (the 7 pt notes, the page 11 capacity line)
  replies.json        the replies page: capacity, knowledge date, entries keyed by question, close, footer
  letter.json         the covering letter paragraphs
  enclosures.json     Section 8 enclosures in pack order, each a PDF under enclosures/
  profile.json        company facts the letterhead needs that the company profile does not hold
  forms/              the requester's blank form (checked by sha256 in manifest.json)
  jobs/<flat>_<date>.json   one release: the flat, the requester, the figures, the dates

Nothing here sends anything, and the signature is never automated: page 11 and the letter are signed by the
director. Every {{placeholder}} must resolve, or rendering stops and names what is missing.
"""
from __future__ import annotations

import json
import re
import shutil
import subprocess
import tempfile
from datetime import date
from pathlib import Path
from typing import Any

from .rtm_templates import BRANDING, render_email

_PH = re.compile(r"\{\{\s*([A-Za-z0-9_.]+)\s*\}\}")
_MONTHS = ["January", "February", "March", "April", "May", "June", "July", "August", "September", "October",
           "November", "December"]

# Page stamp (spec R21): "Page N of M", Helvetica 9, black, right edge 20 pt in from the page edge, baseline
# 12 pt up. Measured from the Flat 7 pack.
STAMP_FONT, STAMP_SIZE, STAMP_RIGHT_INSET, STAMP_BASELINE = "Helvetica", 9, 20.0, 12.0


class TemplateError(ValueError):
    """A template could not be rendered as asked (missing value, missing asset, failed check)."""


# ---------------------------------------------------------------------------------------------- assets

def template_dir() -> Path:
    from ..settings import get_settings
    return Path(get_settings().data_dir) / "rtm_templates" / "lpe1"


def _load(name: str) -> Any:
    p = template_dir() / name
    if not p.exists():
        raise TemplateError(f"LPE1 template asset missing: {p}")
    return json.loads(p.read_text(encoding="utf-8"))


def long_date(iso: str) -> str:
    d = date.fromisoformat(str(iso)[:10])
    return f"{d.day} {_MONTHS[d.month - 1]} {d.year}"


def fill(text: str, ctx: dict[str, Any], *, where: str = "") -> str:
    """Replace every {{name}} from ctx. A missing or blank value is an error, never a silent hole."""
    missing = [m for m in _PH.findall(text or "") if str(ctx.get(m, "")).strip() == ""]
    if missing:
        raise TemplateError(f"{where or 'template'}: no value for {', '.join(sorted(set(missing)))}")
    return _PH.sub(lambda m: str(ctx[m.group(1)]), text or "")


def _pages(pdf: Path) -> int:
    import pypdf
    return len(pypdf.PdfReader(str(pdf)).pages)


def _join(items: list[str]) -> str:
    return items[0] if len(items) == 1 else ", ".join(items[:-1]) + " and " + items[-1]


def _span(first: int, last: int, *, word: str = "page") -> str:
    return f"{word} {first}" if first == last else f"{word}s {first} to {last}"


# ---------------------------------------------------------------------------------------------- context

def build_context(job: dict[str, Any], *, replies_pages: int = 1) -> dict[str, Any]:
    """Everything the templates can refer to: company identity, profile, the job, and derived values
    (long dates, date parts, enclosure page numbers and the sentences that list them)."""
    ctx: dict[str, Any] = {k: v for k, v in BRANDING.items()}
    ctx.update(_load("profile.json"))
    ctx.update(job)
    ctx["release_date_long"] = long_date(job["release_date"])
    ctx["rtm_acquired_long"] = long_date(job["rtm_acquired"])
    for side in ("from", "to"):
        d = date.fromisoformat(job[f"demand_{side}"])
        ctx[f"demand_{side}_dd"], ctx[f"demand_{side}_mm"], ctx[f"demand_{side}_yyyy"] = (
            f"{d.day:02d}", f"{d.month:02d}", str(d.year))
    form = template_dir() / _load("form_answers.json")["form"]
    ctx["form_pages"] = _pages(form)
    ctx["replies_page"] = ctx["form_pages"] + 1
    page = ctx["form_pages"] + replies_pages + 1
    letter_bits, reply_bits = [], []
    for item in _load("enclosures.json")["items"]:
        n = _pages(template_dir() / item["file"])
        first, last = page, page + n - 1
        ctx[f"page_{item['question']}"] = _span(first, last).split(" ", 1)[1] if first == last else f"{first} to {last}"
        letter_bits.append(f"{item['letter']} ({item['question']}, {_span(first, last)})")
        reply_bits.append(f"{item['replies']} ({item['question']}), {_span(first, last)}")
        page = last + 1
    ctx["pack_pages"] = page - 1
    ctx["enclosures_letter"] = _join(letter_bits)
    ctx["enclosures_replies"] = "; ".join(reply_bits) + "."
    return ctx


# ---------------------------------------------------------------------------------------------- form

def _fill_form(ctx: dict[str, Any], out: Path) -> dict[str, Any]:
    """Fill the requester's own form (never a substitute edition), then read every value back."""
    import hashlib
    import pypdf
    from pypdf.generic import NameObject
    answers = _load("form_answers.json")
    src = template_dir() / answers["form"]
    want_sha = (_load("manifest.json").get("form") or {}).get("sha256")
    got_sha = hashlib.sha256(src.read_bytes()).hexdigest()
    if want_sha and got_sha != want_sha:
        raise TemplateError(f"form {src.name} is not the requester's edition (sha256 {got_sha[:16]})")
    text = {k: fill(v, ctx, where=f"form field {k}") for k, v in answers["text"].items()}
    choice = dict(answers["choice"])
    w = pypdf.PdfWriter(clone_from=str(src))
    for page in w.pages:
        w.update_page_form_field_values(page, text, auto_regenerate=False)
    # Radios and checkboxes: set the chosen kid on, every other kid off, and the field value on the parent.
    for page in w.pages:
        for a in page.get("/Annots") or []:
            ao = a.get_object()
            if ao.get("/Subtype") != "/Widget":
                continue
            parent = ao.get("/Parent").get_object() if ao.get("/Parent") else None
            parts, o = [], ao
            while o is not None:
                if o.get("/T") is not None:
                    parts.append(str(o["/T"]))
                o = o.get("/Parent").get_object() if o.get("/Parent") else None
            qname = ".".join(reversed(parts))
            ft = ao.get("/FT") or (parent.get("/FT") if parent else None)
            if ft != "/Btn":
                continue
            states = [k for k in (ao.get("/AP", {}).get("/N", {}) or {}).keys() if k != "/Off"]
            want = choice.get(qname)
            on = bool(want) and want in states
            ao[NameObject("/AS")] = NameObject(want if on else "/Off")
            holder = parent if (parent is not None and ao.get("/T") is None) else ao
            if on:
                holder[NameObject("/V")] = NameObject(want)
    # Fields the form marks multiline that the release draws on one line (page 11 Print Name).
    single = set(answers.get("single_line") or [])
    nudge = {k.split(".")[-1]: v for k, v in (answers.get("nudge") or {}).items()}
    if single:
        for page in w.pages:
            for a in page.get("/Annots") or []:
                ao = a.get_object()
                holder = ao if ao.get("/T") is not None else (ao.get("/Parent").get_object() if ao.get("/Parent") else ao)
                parts, o = [], holder
                while o is not None:
                    if o.get("/T") is not None:
                        parts.append(str(o["/T"]))
                    o = o.get("/Parent").get_object() if o.get("/Parent") else None
                if ".".join(reversed(parts)) in single and holder.get("/Ff") is not None:
                    holder[NameObject("/Ff")] = pypdf.generic.NumberObject(int(holder["/Ff"]) & ~4096)
    # Let poppler draw the text fields itself (it word-wraps multiline fields; pypdf's streams do not).
    w._root_object["/AcroForm"][NameObject("/NeedAppearances")] = pypdf.generic.BooleanObject(True)
    for page in w.pages:
        for a in page.get("/Annots") or []:
            ao = a.get_object()
            ft = ao.get("/FT") or (ao.get("/Parent").get_object().get("/FT") if ao.get("/Parent") else None)
            if ft == "/Tx" and "/AP" in ao:
                del ao["/AP"]
            # Measured placement adjustments: a field's text drawn at [dx, dy] points from where poppler
            # would put it, so the release matches the approved Flat 7 pack exactly.
            name = ao.get("/T") or (ao.get("/Parent").get_object().get("/T") if ao.get("/Parent") else None)
            dxy = nudge.get(str(name))
            if ft == "/Tx" and dxy:
                r = [float(v) for v in ao["/Rect"]]
                ao[NameObject("/Rect")] = pypdf.generic.ArrayObject([pypdf.generic.FloatObject(v) for v in
                                                                     (r[0] + dxy[0], r[1] + dxy[1], r[2] + dxy[0], r[3] + dxy[1])])
    with open(out, "wb") as fh:
        w.write(fh)
    # Readback: every intended value is set, and nothing else is.
    r = pypdf.PdfReader(str(out))
    fields = r.get_fields() or {}
    bad = [k for k, v in text.items() if str((fields.get(k) or {}).get("/V", "")) != v]
    stray = [k for k, f in fields.items() if f.get("/FT") == "/Btn" and str(f.get("/V") or "/Off") not in ("/Off", "")
             and k not in choice]
    wrong = [k for k, v in choice.items() if str((fields.get(k) or {}).get("/V", "")) != v]
    if bad or stray or wrong:
        raise TemplateError(f"form readback failed: text {bad[:5]}, unexpected ticks {stray[:5]}, wrong ticks {wrong[:5]}")
    return {"fields": len(text), "choices": len(choice)}


def _flatten(filled: Path, out: Path) -> None:
    """Flatten by rendering through poppler (spec 6.4); afterwards no form fields may remain."""
    import pypdf
    exe = shutil.which("pdftocairo") or "/opt/homebrew/bin/pdftocairo"
    if not Path(exe).exists():
        raise TemplateError("pdftocairo (poppler) is not installed: brew install poppler")
    subprocess.run([exe, "-pdf", str(filled), str(out)], check=True, capture_output=True, timeout=120)
    if pypdf.PdfReader(str(out)).get_fields():
        raise TemplateError("flattening left form fields in the PDF")


def _overlay(pdf: Path, items: list[dict[str, Any]], out: Path) -> None:
    import io
    import pypdf
    from reportlab.pdfgen import canvas
    r = pypdf.PdfReader(str(pdf))
    w = pypdf.PdfWriter()
    for i, page in enumerate(r.pages):
        mine = [o for o in items if o["page"] == i + 1]
        if mine:
            buf = io.BytesIO()
            c = canvas.Canvas(buf, pagesize=(float(page.mediabox.width), float(page.mediabox.height)))
            for o in mine:
                c.setFont(o.get("font", "Helvetica"), float(o.get("size", 7)))
                c.drawString(float(o["x"]), float(o["y"]), o["text"])
            c.save()
            w.add_page(page).merge_page(pypdf.PdfReader(io.BytesIO(buf.getvalue())).pages[0])
        else:
            w.add_page(page)
    with open(out, "wb") as fh:
        w.write(fh)


def stamp_pages(pdf: Path, out: Path) -> int:
    """Stamp "Page N of M" on every page, right-aligned 20 pt in from the right edge (spec R21)."""
    import io
    import pypdf
    from reportlab.pdfbase.pdfmetrics import stringWidth
    from reportlab.pdfgen import canvas
    r = pypdf.PdfReader(str(pdf))
    total = len(r.pages)
    w = pypdf.PdfWriter()
    for i, page in enumerate(r.pages):
        W, H = float(page.mediabox.width), float(page.mediabox.height)
        label = f"Page {i + 1} of {total}"
        buf = io.BytesIO()
        c = canvas.Canvas(buf, pagesize=(W, H))
        c.setFont(STAMP_FONT, STAMP_SIZE)
        c.drawString(W - STAMP_RIGHT_INSET - stringWidth(label, STAMP_FONT, STAMP_SIZE), STAMP_BASELINE, label)
        c.save()
        w.add_page(page).merge_page(pypdf.PdfReader(io.BytesIO(buf.getvalue())).pages[0])
    with open(out, "wb") as fh:
        w.write(fh)
    return total


# ---------------------------------------------------------------------------------------------- docx

def _docx_to_pdf(docx: Path) -> Path:
    exe = (shutil.which("soffice") or "/opt/homebrew/bin/soffice")
    if not Path(exe).exists():
        raise TemplateError("LibreOffice (soffice) is needed to turn the letter and replies into PDF")
    subprocess.run([exe, "--headless", "--convert-to", "pdf", "--outdir", str(docx.parent), str(docx)],
                   check=True, capture_output=True, timeout=180)
    pdf = docx.with_suffix(".pdf")
    if not pdf.exists():
        raise TemplateError(f"LibreOffice produced no PDF for {docx.name}")
    return pdf


def _run(p, text: str, *, size: float, bold: bool = False, italic: bool = False, colour: str | None = None):
    from docx.shared import Pt, RGBColor
    r = p.add_run(text)
    r.font.name = "Calibri"
    r.font.size = Pt(size)
    r.bold, r.italic = bold, italic
    if colour:
        r.font.color.rgb = RGBColor.from_string(colour)
    return r


def _para(doc, *, before: float = 0, after: float = 0, line: float | None = None, left: float = 0,
          hanging: float = 0, border_bottom: tuple[str, float, float] | None = None):
    from docx.enum.text import WD_LINE_SPACING
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before, pf.space_after = Pt(before), Pt(after)
    if line is not None:
        pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        pf.line_spacing = Pt(line)
    if left or hanging:
        pf.left_indent = Pt(left)
        pf.first_line_indent = Pt(-hanging)
        pf.tab_stops.add_tab_stop(Pt(left))
    if border_bottom:
        colour, width_pt, space_pt = border_bottom
        pPr = p._p.get_or_add_pPr()
        bdr = OxmlElement("w:pBdr")
        b = OxmlElement("w:bottom")
        b.set(qn("w:val"), "single")
        b.set(qn("w:sz"), str(int(width_pt * 8)))
        b.set(qn("w:space"), str(int(space_pt)))
        b.set(qn("w:color"), colour)
        bdr.append(b)
        pPr.append(bdr)
    return p


def _new_doc(margin_lr: float, margin_top: float, margin_bottom: float):
    import docx
    from docx.shared import Pt
    d = docx.Document()
    s = d.sections[0]
    s.page_width, s.page_height = Pt(595.3), Pt(841.9)
    s.left_margin = s.right_margin = Pt(margin_lr)
    s.top_margin, s.bottom_margin = Pt(margin_top), Pt(margin_bottom)
    st = d.styles["Normal"]
    st.font.name = "Calibri"
    st.paragraph_format.space_after = Pt(0)
    return d


# Layout measured from the Flat 7 letter (LibreOffice, Carlito): 23 mm side margins, 10.5 pt body.
LETTER = {"margin_lr": 65.1, "margin_top": 56.7, "margin_bottom": 56.7, "body": 10.5, "addr_line": 12.8,
          "body_line": 13.2, "gap": 7.0, "grey": "7F7F7F"}


def render_letter(ctx: dict[str, Any], out_dir: Path) -> Path:
    L, spec = LETTER, _load("letter.json")
    d = _new_doc(L["margin_lr"], L["margin_top"], L["margin_bottom"])
    body, g = L["body"], L["gap"]
    _run(_para(d, line=14.4), ctx["company"], size=12, bold=True)
    _run(_para(d, line=10.4), f"Company No. {ctx['company_number']}   ·   {ctx['email']}   ·   {ctx['tel']}",
         size=8.5, colour=L["grey"])
    _run(_para(d, line=10.4, after=16, border_bottom=(L["grey"], 0.75, 4)), ctx["letterhead_address"], size=8.5)
    for i, t in enumerate(spec["recipient_lines"]):
        _run(_para(d, line=L["addr_line"]), fill(t, ctx, where="letter"), size=body)
    _run(_para(d, line=L["addr_line"], before=12), ctx["release_date_long"], size=body)
    for i, t in enumerate(spec["refs"]):
        _run(_para(d, line=L["addr_line"], before=12 if i == 0 else 0), fill(t, ctx, where="letter"), size=body)
    _run(_para(d, line=L["addr_line"], before=12), spec["salutation"], size=body)
    for i, t in enumerate(spec["heading"]):
        _run(_para(d, line=13.4, before=9.5 if i == 0 else 0), fill(t, ctx, where="letter"), size=body, bold=True)
    after_list = False
    for item in spec["paragraphs"]:
        raw = item["text"]
        if item.get("optional") and not all(str(ctx.get(m, "")).strip() for m in _PH.findall(raw)):
            continue
        text = fill(raw, ctx, where="letter")
        if "item" in item:
            p = _para(d, line=L["body_line"], before=4.0 if after_list else 4.2, left=18, hanging=18)
            _run(p, f"{item['item']}\t{text}", size=body)
        else:
            # A paragraph straight after the numbered list sits closer (measured: 4 pt, not 7).
            p = _para(d, line=L["body_line"], before=item.get("space_before", 4.0 if after_list else g))
            if item.get("lead"):
                _run(p, item["lead"] + "  ", size=body, bold=True)
            _run(p, text, size=body)
        after_list = "item" in item
    _run(_para(d, line=L["body_line"], before=9.4), spec["close"], size=body)
    for i, t in enumerate(spec["signature_lines"]):
        # The gap above the name is the space the director signs in; the signature is never automated.
        _run(_para(d, line=L["addr_line"], before=31.9 if i == 0 else 0), fill(t, ctx, where="letter"),
             size=body, bold=(i == 0))
    docx_path = out_dir / f"Stage2_1_Cover_Letter_{ctx['flat_id']}_LPE1.docx"
    d.save(str(docx_path))
    return _docx_to_pdf(docx_path)


# Layout measured from the Flat 7 replies page: 20 mm margins, 10 pt body, navy title, grey footer.
REPLIES = {"margin_lr": 56.7, "margin_top": 49.9, "margin_bottom": 56.7, "title": 15, "body": 10,
           "line": 12.5, "entry_gap": 3.5, "navy": "1F3864", "grey_note": "595959", "grey_footer": "404040",
           "rule": "BFBFBF"}


def render_replies(ctx: dict[str, Any], out_dir: Path) -> Path:
    from docx.shared import Pt
    R, spec = REPLIES, _load("replies.json")
    d = _new_doc(R["margin_lr"], R["margin_top"], R["margin_bottom"])
    b, ln = R["body"], R["line"]
    _run(_para(d, line=18), fill(spec["title"], ctx, where="replies"), size=R["title"], bold=True, colour=R["navy"])
    _run(_para(d, line=ln, before=3.9), fill(spec["property_line"], ctx, where="replies"), size=b)
    p = _para(d, line=ln, before=6.2)
    _run(p, spec["capacity"]["heading"] + "  ", size=b, bold=True)
    _run(p, fill(spec["capacity"]["text"], ctx, where="replies"), size=b)
    _run(_para(d, line=12.2, before=3.3, border_bottom=(R["rule"], 0.5, 4.2)),
         fill(spec["knowledge"], ctx, where="replies"), size=b, italic=True, colour=R["grey_note"])
    entries = spec["entries"]
    for i, e in enumerate(entries):
        last = i == len(entries) - 1
        p = _para(d, line=ln, before=R["entry_gap"] if i else 8.2,
                  border_bottom=(R["rule"], 0.5, 10.1) if last else None)
        _run(p, e["heading"] + "  ", size=b, bold=True)
        _run(p, fill(e["text"], ctx, where=f"replies {e['heading']}"), size=b)
    _run(_para(d, line=12.2, before=5.8), fill(spec["close"], ctx, where="replies"), size=b, italic=True)
    d.sections[0].footer_distance = Pt(33.6)
    fp = d.sections[0].footer.paragraphs[0]
    fp.paragraph_format.space_after = Pt(0)
    _run(fp, fill(spec["footer"], ctx, where="replies"), size=9, colour=R["grey_footer"])
    docx_path = out_dir / f"LPE1_replies_{ctx['flat_id']}.docx"
    d.save(str(docx_path))
    return _docx_to_pdf(docx_path)


# ---------------------------------------------------------------------------------------------- pack

def render_pack(job: dict[str, Any], out_dir: Path) -> dict[str, Any]:
    """The whole stage 2 release for one job: letter PDF, pack PDF (form, replies, enclosures, stamped),
    and the email draft. Nothing is sent and nothing is signed."""
    import pypdf
    out_dir.mkdir(parents=True, exist_ok=True)
    work = Path(tempfile.mkdtemp(prefix="lpe1_", dir=str(out_dir)))
    ctx = build_context(job)
    replies = render_replies(ctx, work)
    n = _pages(replies)
    if n != 1:  # enclosure page numbers depend on how long the replies run
        ctx = build_context(job, replies_pages=n)
        replies = render_replies(ctx, work)
    filled, flat, form = work / "filled.pdf", work / "flat.pdf", work / "form.pdf"
    readback = _fill_form(ctx, filled)
    _flatten(filled, flat)
    _overlay(flat, [o for o in _load("form_overlays.json")["overlays"]], form)
    w = pypdf.PdfWriter()
    for part in [form, replies] + [template_dir() / i["file"] for i in _load("enclosures.json")["items"]]:
        for page in pypdf.PdfReader(str(part)).pages:
            w.add_page(page)
    joined = work / "joined.pdf"
    with open(joined, "wb") as fh:
        w.write(fh)
    pack = out_dir / f"Stage2_2_LPE1_Pack_{ctx['flat_id']}_SIGN_PAGE_{ctx['form_pages']}.pdf"
    total = stamp_pages(joined, pack)
    if total != ctx["pack_pages"]:
        raise TemplateError(f"pack has {total} pages but the letter and replies say {ctx['pack_pages']}")
    letter = render_letter(ctx, work)
    letter_out = out_dir / f"Stage2_1_Cover_Letter_{ctx['flat_id']}_LPE1.pdf"
    shutil.copyfile(letter, letter_out)
    email = render_email("cover_email_lpe1_stage2", {k: ctx[k] for k in
                         ("recipient_email", "email_salutation", "property_short", "their_ref")})
    shutil.rmtree(work, ignore_errors=True)
    return {"ok": True, "letter_pdf": str(letter_out), "pack_pdf": str(pack), "pack_pages": total,
            "email": email, "form_readback": readback,
            "to_sign": [f"{pack.name}, page {ctx['form_pages']}", letter_out.name]}


def render_job(job_file: str, out_dir: str | Path | None = None) -> dict[str, Any]:
    job = json.loads((template_dir() / "jobs" / job_file).read_text(encoding="utf-8"))
    if out_dir is None:
        from ..settings import get_settings
        out_dir = Path(get_settings().data_dir) / "rtm_generated" / "lpe1" / job["flat_id"]
    return render_pack(job, Path(out_dir))


# ---------------------------------------------------------------------------------------------- proof

def compare_pdfs(generated: str | Path, reference: str | Path, *, tolerance_pt: float = 1.0) -> dict[str, Any]:
    """Character-level comparison of two PDFs, page by page. Each reference character is matched to the
    nearest identical character on the same page of the generated PDF; a page passes when every character
    on both sides is matched within tolerance_pt. Reports the largest move and anything unmatched."""
    import pdfplumber  # installed alongside poppler for this check
    out: dict[str, Any] = {"pages": []}
    with pdfplumber.open(str(generated)) as g, pdfplumber.open(str(reference)) as r:
        out["page_count"] = [len(g.pages), len(r.pages)]
        out["same_page_count"] = len(g.pages) == len(r.pages)
        for i, (pg, pr) in enumerate(zip(g.pages, r.pages)):
            gen = [c for c in pg.chars if c["text"].strip()]
            used = [False] * len(gen)
            by_text: dict[str, list[int]] = {}
            for k, c in enumerate(gen):
                by_text.setdefault(c["text"], []).append(k)
            worst, missing = 0.0, []
            for c in (c for c in pr.chars if c["text"].strip()):
                best, bi = None, None
                for k in by_text.get(c["text"], []):
                    if used[k]:
                        continue
                    d = max(abs(gen[k]["x0"] - c["x0"]), abs(gen[k]["top"] - c["top"]))
                    if best is None or d < best:
                        best, bi = d, k
                if best is None or best > tolerance_pt:
                    missing.append(c["text"])
                    continue
                used[bi] = True
                worst = max(worst, best)
            extra = [gen[k]["text"] for k in range(len(gen)) if not used[k]]
            out["pages"].append({"page": i + 1, "ok": not missing and not extra, "max_move": round(worst, 3),
                                 "missing": "".join(missing)[:80], "extra": "".join(extra)[:80]})
    out["ok"] = out["same_page_count"] and all(p["ok"] for p in out["pages"])
    return out


def acceptance_check(job_file: str | None = None, tolerance_pt: float = 1.0) -> dict[str, Any]:
    """Render a job whose release was approved and compare it with the PDFs that were sent (kept under
    reference/). Passes when every page has the same text, with no character more than tolerance_pt away.
    The default job is the first live release: Flat 7 of the site profile's first block, 23 Sep 2026."""
    if job_file is None:
        from .. import profile as _profile
        job_file = f"{_profile.block_codes()[0]}F7_2026-09-23.json"
    ref = template_dir() / "reference"
    packs = sorted(ref.glob("Stage2_2_LPE1_Pack_*.pdf"))
    letters = sorted(ref.glob("Stage2_1_Cover_Letter_*.pdf"))
    if not packs or not letters:
        raise TemplateError(f"no reference PDFs under {ref}")
    with tempfile.TemporaryDirectory() as tmp:
        res = render_job(job_file, tmp)
        report = {"pack": compare_pdfs(res["pack_pdf"], packs[0], tolerance_pt=tolerance_pt),
                  "letter": compare_pdfs(res["letter_pdf"], letters[0], tolerance_pt=tolerance_pt)}
    report["ok"] = report["pack"]["ok"] and report["letter"]["ok"]
    return report
